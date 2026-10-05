"""Validate facilitator/class-N.md and copy each slide's notes into the deck's presenter notes.

Usage: python notes_to_pptx.py <decks_dir> [--check]
  decks_dir: folder holding "Class N - .../IJ#Nfinal.pptx" (the frontmatter `deck` path is relative to it).
Writes IJ#N-notes.pptx next to each original. Originals are never modified.
Requires python-pptx.
"""
import re
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent.parent  # facilitator/
SECTIONS = ["Teacher Intro", "Class Intro", "Head: What & Why", "Hand: Motivation", "Heart: Tools", "Group Time Handoff"]


def plain(md: str) -> str:
    """Markdown bullets -> presenter-view text."""
    out = []
    for line in md.strip().splitlines():
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
        line = re.sub(r"(?<!\w)[*_](.+?)[*_](?!\w)", r"\1", line)
        line = re.sub(r"^(\s*)- ", lambda m: m.group(1) + "• ", line)
        out.append(line)
    return "\n".join(out)


def parse(path: Path):
    text = path.read_text()
    fm = re.match(r"---\n(.*?)\n---\n", text, re.S)
    if not fm:
        raise ValueError(f"{path.name}: missing frontmatter")
    meta = dict(re.findall(r'^(\w+):\s*"?(.*?)"?\s*$', fm.group(1), re.M))
    sections, slides, errors = [], {}, []
    current = None
    for chunk in re.split(r"(?m)^(?=## |### Slide )", text[fm.end():]):
        head = chunk.splitlines()[0] if chunk.strip() else ""
        if m := re.match(r"## (.+?) \((\d+) min\)", head):
            current = (m.group(1), int(m.group(2)))
            sections.append(current)
        elif m := re.match(r"### Slide (\d+) — (.*)", head):
            n = int(m.group(1))
            if n in slides:
                errors.append(f"slide {n} duplicated")
            body = chunk.split("\n", 1)[1] if "\n" in chunk else ""
            prefix = ""
            if current and not any(v[2] == current for v in slides.values()):
                prefix = f"[{current[0]} · {current[1]} min]\n"
            slides[n] = (m.group(2), prefix + plain(body), current)
    names = [s[0] for s in sections]
    if names != SECTIONS:
        errors.append(f"sections {names} != {SECTIONS}")
    total = sum(s[1] for s in sections)
    if total != 60:
        errors.append(f"minutes total {total}, expected 60")
    if list(slides) != sorted(slides):
        errors.append("slides out of deck order")
    return meta, slides, errors


def main():
    decks = Path(sys.argv[1]).expanduser()
    check_only = "--check" in sys.argv
    failed = False
    for md in sorted(HERE.glob("class-*.md")):
        meta, slides, errors = parse(md)
        deck_path = decks / meta["deck"]
        prs = Presentation(deck_path)
        count = len(prs.slides)
        if sorted(slides) != list(range(1, count + 1)):
            missing = sorted(set(range(1, count + 1)) - set(slides))
            extra = sorted(set(slides) - set(range(1, count + 1)))
            errors.append(f"deck has {count} slides; missing {missing}, extra {extra}")
        reviews = md.read_text().count("⚑ Review:")
        status = "OK" if not errors else "FAIL: " + "; ".join(errors)
        print(f"{md.name}: {len(slides)}/{count} slides, {reviews} review flags — {status}")
        if errors:
            failed = True
            continue
        if check_only:
            continue
        for i, slide in enumerate(prs.slides, start=1):
            slide.notes_slide.notes_text_frame.text = slides[i][1]
        out = deck_path.with_name(deck_path.stem.replace("final", "") + "-notes.pptx")
        prs.save(out)
        print(f"  wrote {out}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
